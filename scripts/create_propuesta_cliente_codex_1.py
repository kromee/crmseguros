from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs" / "PROPUESTA-CLIENTE-CODEX-1.docx"


def set_font(run, name: str, size: int, bold: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 0, line: float = 1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_fixed_width(table, widths_in):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_in) * 1440)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def style_body_paragraph(paragraph, *, after: int = 8, line: float = 1.333) -> None:
    set_paragraph_spacing(paragraph, before=0, after=after, line=line)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        set_font(run, "Calibri", 11, color="000000")


def add_heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    if level == 1:
        p.style = "Heading 1"
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_font(r, "Calibri", 16, bold=False, color="2E74B5")
    elif level == 2:
        p.style = "Heading 2"
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_font(r, "Calibri", 13, bold=False, color="2E74B5")
    else:
        p.style = "Heading 3"
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_font(r, "Calibri", 12, bold=False, color="1F4D78")


def add_body(doc, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    style_body_paragraph(p)
    r = p.add_run(text)
    set_font(r, "Calibri", 11, bold=bold, color="000000")
    r.italic = italic


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        r = p.add_run(item)
        set_font(r, "Calibri", 11, color="000000")


def add_scope_table(doc) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed_width(table, [2.1, 4.4])
    table.style = "Table Grid"

    headers = ["Área", "Funcionalidad"]
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = ""
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, "Calibri", 10, bold=True, color="1F3A5F")

    rows = [
        ("Contactos", "Directorio, alta guiada, detalle, foto, asignación por agente"),
        ("Pólizas", "Vida y auto, vencimientos, renovaciones, archivos PDF"),
        ("Pensiones", "Asesoría y trámites (IMSS Ley 73 / AFORE Ley 97)"),
        ("Trámites vehiculares", "Alta, baja, placas, cotización, checklist documental"),
        ("Pipeline", "Tablero Kanban por etapas, prioridad, probabilidad de cierre"),
        ("Calendario", "Eventos, recordatorios y alertas operativas"),
        ("Bitácora", "Historial de llamadas, WhatsApp, visitas y notas"),
        ("Finanzas", "Pagos vinculados a pólizas, resumen de cobranza"),
        ("Panel de control", "Indicadores de cartera, renovaciones y alertas"),
        ("Usuarios", "Roles administrador y operativo"),
    ]

    for area, functionality in rows:
        row = table.add_row().cells
        left = row[0]
        right = row[1]
        left.text = ""
        right.text = ""
        for cell, value, align in ((left, area, WD_ALIGN_PARAGRAPH.LEFT), (right, functionality, WD_ALIGN_PARAGRAPH.LEFT)):
            p = cell.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_font(r, "Calibri", 10, color="000000")



def add_key_value_table(doc) -> None:
    table = doc.add_table(rows=3, cols=2)
    set_table_fixed_width(table, [1.75, 4.75])
    table.style = "Table Grid"

    labels = ["Cliente", "Fecha", "Versión"]
    values = [
        "Propuesta comercial",
        date.today().strftime("%d/%m/%Y"),
        "1.1",
    ]

    for row_idx, (label, value) in enumerate(zip(labels, values, strict=True)):
        left = table.cell(row_idx, 0)
        right = table.cell(row_idx, 1)
        left.text = ""
        right.text = ""
        set_cell_shading(left, "F2F4F7")
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.space_before = Pt(0)
        lp.paragraph_format.space_after = Pt(0)
        rp.paragraph_format.space_before = Pt(0)
        rp.paragraph_format.space_after = Pt(0)

        lrun = lp.add_run(label)
        set_font(lrun, "Calibri", 10, bold=True, color="1F3A5F")
        rrun = rp.add_run(value)
        set_font(rrun, "Calibri", 10, color="000000")


def add_pricing_table(doc) -> None:
    table = doc.add_table(rows=5, cols=2)
    set_table_fixed_width(table, [2.55, 4.05])
    table.style = "Table Grid"

    headers = ["Concepto", "Importe"]
    rows = [
        ("Implementación inicial", "$25,000 MXN"),
        ("Mensualidad del sistema", "$1,500 MXN"),
        ("Hosting / infraestructura", "Incluido"),
        ("Total mensual estimado", "$1,500 MXN"),
    ]

    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = ""
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, "Calibri", 10, bold=True, color="1F3A5F")

    for row_idx, (concept, value) in enumerate(rows, start=1):
        for col_idx, txt in enumerate((concept, value)):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            set_font(r, "Calibri", 10, bold=(row_idx == 4), color="000000")
            if row_idx == 4:
                set_cell_shading(cell, "F4F6F9")


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("000000")

    # Title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("PROPUESTA COMERCIAL")
    set_font(run, "Calibri", 24, bold=True, color="0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("CRM SaaS para agencia de seguros con 2 usuarios activos")
    set_font(run, "Calibri", 11, bold=False, color="555555")

    add_key_value_table(doc)

    add_body(
        doc,
        "Esta propuesta presenta una implementación inicial y una renta mensual pensada para una agencia pequeña que "
        "necesita un CRM a medida, con acceso para 2 usuarios, autenticación, carga de archivos, hosting incluido y despliegue administrado.",
    )

    add_heading(doc, "1. Alcance del sistema", level=1)
    add_body(doc, "El servicio considera la plataforma lista para operar con los módulos principales:")
    add_bullets(
        doc,
        [
            "Gestión de contactos, prospectos, pólizas, pagos y calendario.",
            "Control de acceso por roles y operación multiusuario.",
            "Carga y consulta de archivos relacionados con clientes y pólizas.",
            "Hosting básico incluido dentro de la mensualidad.",
            "Ajustes y soporte evolutivo menor para la operación diaria.",
        ],
    )

    add_heading(doc, "2. Inversión propuesta", level=1)
    add_body(
        doc,
        "La propuesta está pensada para ser accesible en una primera etapa, sin perder margen operativo "
        "para soporte, mantenimiento, hosting y mejoras menores.",
    )
    add_pricing_table(doc)

    add_body(
        doc,
        "La implementación inicial cubre la configuración base, puesta en producción, validación del entorno, "
        "personalización básica y ajustes de arranque.",
    )

    add_heading(doc, "3. Condiciones comerciales", level=1)
    add_bullets(
        doc,
        [
            "Pago mensual anticipado.",
            "Se puede contratar en modalidad anual (12 meses) o a 24 meses, con precio congelado durante la vigencia del contrato.",
            "La mensualidad incluye hosting básico administrado, soporte y mantenimiento sobre lo ya desarrollado.",
        "Agregar usuarios adicionales tiene costo y se cotiza por separado.",
            "En contrato de 24 meses se otorga prioridad en soporte sobre lo ya desarrollado.",
            "No incluye dominio, correo corporativo ni integraciones de terceros no contempladas al arranque.",
            "Nuevos desarrollos o funcionalidades adicionales se cotizan por separado.",
            "Vigencia de la propuesta: 15 días naturales.",
        ],
    )

    add_heading(doc, "4. Cierre", level=1)
    add_body(
        doc,
        "Si esta propuesta te hace sentido, el siguiente paso es confirmar el alcance inicial, elegir la modalidad de contratación y programar la salida a producción. A partir de ahí se puede afinar cualquier detalle operativo o de crecimiento.",
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(DOCX_PATH)
    print(f"Created {DOCX_PATH}")


if __name__ == "__main__":
    main()
